from pathlib import Path

from pptx import Presentation
from langchain_core.documents import Document


def load_powerpoint(file_path: Path):

    presentation = Presentation(file_path)

    documents = []

    # =====================================================
    # READ COURSE AND WEEK FROM FOLDER STRUCTURE
    # =====================================================

    course = file_path.parent.parent.name
    week = file_path.parent.name


    # =====================================================
    # LOAD SLIDES
    # =====================================================

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:

                    slide_text.append(text)


        full_text = "\n".join(slide_text)


        if full_text:

            document = Document(

                page_content=full_text,

                metadata={

                    "source": file_path.name,

                    "course": course,

                    "week": week,

                    "slide": slide_number,

                    "document_type": "lecture"
                }
            )

            documents.append(document)


    return documents