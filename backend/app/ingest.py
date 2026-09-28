from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader
import docx2txt

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_chroma import Chroma

from embeddings import get_embeddings


# ===================================
# PROJECT PATHS
# ===================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]

LECTURES_DIR = (
    PROJECT_ROOT
    / "data"
    / "lectures"
)

CHROMA_DIR = (
    PROJECT_ROOT
    / "data"
    / "chroma_db"
)


# ===================================
# FIND ALL SUPPORTED FILES
# ===================================

supported_extensions = {
    ".pptx",
    ".pdf",
    ".docx"
}

files = [
    file_path
    for file_path in LECTURES_DIR.rglob("*")
    if file_path.is_file()
    and file_path.suffix.lower()
    in supported_extensions
]


print("\n===================================")
print("      AI UNIVERSITY INGESTION")
print("===================================")

print(
    "\nSupported files found:",
    len(files)
)


# ===================================
# LOAD DOCUMENTS
# ===================================

documents = []


for file_path in files:

    print("\n-----------------------------------")
    print("Loading:", file_path.name)
    print("-----------------------------------")

    course = file_path.parent.parent.name
    week = file_path.parent.name

    print("Course:", course)
    print("Week:", week)
    print("Type:", file_path.suffix.lower())


    # ===================================
    # POWERPOINT
    # ===================================

    if file_path.suffix.lower() == ".pptx":

        presentation = Presentation(file_path)

        slide_count = 0

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            slide_text = []

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:
                        slide_text.append(text)

            full_text = "\n".join(slide_text)

            if full_text:

                document = Document(
                    page_content=full_text,
                    metadata={
                        "source": file_path.name,
                        "course": course,
                        "week": week,
                        "slide": slide_number,
                        "document_type": "powerpoint"
                    }
                )

                documents.append(document)

                slide_count += 1

        print(
            "Slides with text:",
            slide_count
        )


    # ===================================
    # PDF
    # ===================================

    elif file_path.suffix.lower() == ".pdf":

        reader = PdfReader(file_path)

        page_count = 0

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            text = page.extract_text()

            if text:

                text = text.strip()

            if text:

                document = Document(
                    page_content=text,
                    metadata={
                        "source": file_path.name,
                        "course": course,
                        "week": week,
                        "page": page_number,
                        "document_type": "pdf"
                    }
                )

                documents.append(document)

                page_count += 1

        print(
            "PDF pages with text:",
            page_count
        )


    # ===================================
    # WORD DOCUMENT
    # ===================================

    elif file_path.suffix.lower() == ".docx":

        text = docx2txt.process(
            str(file_path)
        )

        text = text.strip()

        if text:

            document = Document(
                page_content=text,
                metadata={
                    "source": file_path.name,
                    "course": course,
                    "week": week,
                    "document_type": "word"
                }
            )

            documents.append(document)

            print(
                "Word document loaded."
            )

        else:

            print(
                "No text found in Word document."
            )


# ===================================
# DOCUMENT SUMMARY
# ===================================

print("\n===================================")
print("DOCUMENTS LOADED")
print("===================================")

print(
    "Total documents:",
    len(documents)
)


# ===================================
# CHUNKING
# ===================================

print("\nCreating chunks...")

splitter = RecursiveCharacterTextSplitter(
    chunk_size=500,
    chunk_overlap=100
)

chunks = splitter.split_documents(
    documents
)

print(
    "Total chunks:",
    len(chunks)
)


# ===================================
# EMBEDDINGS
# ===================================

print("\nLoading embedding model...")

embeddings = get_embeddings()


# ===================================
# CHROMA
# ===================================

print("\nConnecting to Chroma...")

old_store = Chroma(
    persist_directory=str(CHROMA_DIR),
    collection_name="university_material",
    embedding_function=embeddings
)

print(
    "Deleting old Chroma collection..."
)

old_store.delete_collection()

print(
    "Old collection deleted."
)


print(
    "\nCreating fresh Chroma database..."
)

vector_store = Chroma.from_documents(
    documents=chunks,
    embedding=embeddings,
    persist_directory=str(CHROMA_DIR),
    collection_name="university_material"
)


# ===================================
# COMPLETE
# ===================================

print("\n===================================")
print("      INGESTION COMPLETE")
print("===================================")

print(
    "Files:",
    len(files)
)

print(
    "Documents:",
    len(documents)
)

print(
    "Chunks:",
    len(chunks)
)

print(
    "Chroma database:",
    CHROMA_DIR
)


# ===================================
# METADATA CHECK
# ===================================

print("\n===================================")
print("      METADATA CHECK")
print("===================================")


for chunk in chunks[:10]:

    print(
        "\nCourse:",
        chunk.metadata.get("course")
    )

    print(
        "Week:",
        chunk.metadata.get("week")
    )

    print(
        "Source:",
        chunk.metadata.get("source")
    )

    print(
        "Type:",
        chunk.metadata.get("document_type")
    )

    if "slide" in chunk.metadata:

        print(
            "Slide:",
            chunk.metadata.get("slide")
        )

    if "page" in chunk.metadata:

        print(
            "Page:",
            chunk.metadata.get("page")
        )